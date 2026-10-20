"""Genera la plantilla de Excel que se sube a Google Sheets (Archivo > Importar)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Agenda_Dr_Rivas_PLANTILLA.xlsx')
BLACK = '0A0A0A'
head_fill = PatternFill('solid', fgColor=BLACK)
head_font = Font(bold=True, color='FFFFFF')
soft = PatternFill('solid', fgColor='EFECE5')
thin = Side(style='thin', color='DDDDDD')

wb = Workbook()

# ---------------------------------------------------------------- Citas
ws = wb.active
ws.title = 'Citas'
cols = [
    ('ID', 10), ('Registrada', 19), ('Sede', 12), ('Fecha de la cita', 16), ('Hora', 9),
    ('Tipo de cita', 22), ('Nombre', 28), ('Celular', 14), ('Correo', 28), ('Estado', 14),
    ('Notas (uso interno)', 30), ('ID evento calendario', 24), ('Recordatorio enviado', 20),
]
for i, (name, w) in enumerate(cols, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.fill, c.font = head_fill, head_font
    c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.column_dimensions[chr(64 + i)].width = w
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:M1000'
for r in range(2, 1001):
    ws.cell(row=r, column=2).number_format = 'dd/mm/yyyy hh:mm'
    ws.cell(row=r, column=4).number_format = 'dd/mm/yyyy'
    ws.cell(row=r, column=5).number_format = '@'   # Hora como texto (08:30)
    ws.cell(row=r, column=8).number_format = '@'   # Celular como texto

dv_estado = DataValidation(type='list', formula1='"Pendiente,Confirmada,Atendida,No asistió,Cancelada"', allow_blank=True)
dv_sede = DataValidation(type='list', formula1='"Yopal,Aguazul,Bogotá"', allow_blank=True)
ws.add_data_validation(dv_estado); ws.add_data_validation(dv_sede)
dv_estado.add('J2:J1000'); dv_sede.add('C2:C1000')

ws.conditional_formatting.add('A2:M1000', FormulaRule(formula=['$J2="Cancelada"'], font=Font(color='9A9A9A', strike=True)))
ws.conditional_formatting.add('A2:M1000', FormulaRule(formula=['$J2="Confirmada"'], fill=PatternFill('solid', bgColor='E8F3EA', fgColor='E8F3EA')))
ws.conditional_formatting.add('A2:M1000', FormulaRule(formula=['$J2="Pendiente"'], fill=PatternFill('solid', bgColor='FFF7E0', fgColor='FFF7E0')))

# ---------------------------------------------------------------- Configuración
cf = wb.create_sheet('Configuración')
cf.column_dimensions['A'].width = 44
for col, w in zip('BCDEF', (24, 34, 18, 34, 46)):
    cf.column_dimensions[col].width = w

def title(cell, text):
    cf[cell] = text
    cf[cell].font = Font(bold=True, size=13)

title('A1', 'CONFIGURACIÓN DE LA AGENDA  (edita solo las celdas con fondo gris)')
cf['A2'] = 'Los cambios se reflejan en el sitio en pocos minutos. No cambies el orden de filas ni columnas.'
cf['A2'].font = Font(italic=True, color='666666')

title('A3', 'Parámetros')
params = [
    ('Días de anticipación máximos para agendar', 30),
    ('Horas mínimas de anticipación', 2),
    ('Máximo de citas por día en cada sede', 20),
    ('Duración de cada cita (minutos)', 30),
    ('Nombre que ven los pacientes en los correos', 'Clínica Dr. Rivas'),
    ('Correo interno que recibe copia de todas las citas (opcional)', ''),
]
for i, (k, v) in enumerate(params, start=4):
    cf.cell(row=i, column=1, value=k)
    c = cf.cell(row=i, column=2, value=v); c.fill = soft; c.border = Border(bottom=thin)

title('A11', 'Sedes')
for j, h in enumerate(['Clave (no cambiar)', 'Nombre', 'Correo que recibe los avisos de la sede', 'WhatsApp', 'ID del calendario de la sede (vacío = calendario principal)', 'Dirección'], start=1):
    c = cf.cell(row=12, column=j, value=h); c.fill, c.font = head_fill, head_font
    c.alignment = Alignment(wrap_text=True, vertical='center')
cf.row_dimensions[12].height = 44
sedes = [
    ('yopal', 'Yopal', '[correo-yopal@dominio.com]', '573134693232', '', 'Carrera 24 con Calle 12, esquina diagonal al Edificio Luma'),
    ('aguazul', 'Aguazul', '[correo-aguazul@dominio.com]', '573207451448', '', 'Carrera 18A #10-24'),
    ('bogota', 'Bogotá', '[correo-bogota@dominio.com]', '573207671072', '', 'Autopista Norte con Calle 106'),
]
for i, row in enumerate(sedes, start=13):
    for j, v in enumerate(row, start=1):
        c = cf.cell(row=i, column=j, value=v)
        c.border = Border(bottom=thin)
        if j in (3, 5):
            c.fill = soft

title('A18', 'Horarios de atención por sede (formato 24 h, ej. 08:00 y 18:00; vacío = cerrado)')
for j, h in enumerate(['Sede (no cambiar)', 'Día (no cambiar)', 'Desde', 'Hasta'], start=1):
    c = cf.cell(row=19, column=j, value=h); c.fill, c.font = head_fill, head_font
dias = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
r = 20
for key in ('yopal', 'aguazul', 'bogota'):
    for d_i, d in enumerate(dias):
        cf.cell(row=r, column=1, value=key)
        cf.cell(row=r, column=2, value=d)
        desde, hasta = ('', '')
        if d_i <= 4: desde, hasta = '08:00', '18:00'
        elif d_i == 5: desde, hasta = '08:00', '12:00'
        for col, v in ((3, desde), (4, hasta)):
            c = cf.cell(row=r, column=col, value=v)
            c.fill = soft; c.number_format = '@'; c.alignment = Alignment(horizontal='center')
        r += 1
cf['F19'] = 'Nota: los horarios de ejemplo (L-V 8:00-18:00, sábados 8:00-12:00) son provisionales; confírmalos con cada sede.'
cf['F19'].font = Font(italic=True, color='B3261E')
cf['F19'].alignment = Alignment(wrap_text=True, vertical='top')

# ---------------------------------------------------------------- Resumen
rs = wb.create_sheet('Resumen')
rs.column_dimensions['A'].width = 26
for col in 'BCDEF':
    rs.column_dimensions[col].width = 14
rs['A1'] = 'RESUMEN DE CITAS'; rs['A1'].font = Font(bold=True, size=13)
estados = ['Pendiente', 'Confirmada', 'Atendida', 'No asistió', 'Cancelada']
rs['A3'] = 'Sede'
for j, e in enumerate(estados, start=2):
    rs.cell(row=3, column=j, value=e)
for c in rs[3]:
    if c.value:
        c.fill, c.font = head_fill, head_font
for i, s in enumerate(['Yopal', 'Aguazul', 'Bogotá'], start=4):
    rs.cell(row=i, column=1, value=s)
    for j, e in enumerate(estados, start=2):
        col = chr(64 + j)
        rs.cell(row=i, column=j, value=f'=COUNTIFS(Citas!$C$2:$C$1000,$A{i},Citas!$J$2:$J$1000,{col}$3)')
rs['A8'] = 'Citas de hoy'; rs['A8'].font = Font(bold=True)
rs['B8'] = '=COUNTIFS(Citas!$D$2:$D$1000,TODAY(),Citas!$J$2:$J$1000,"<>Cancelada")'
rs['A9'] = 'Citas de mañana'; rs['A9'].font = Font(bold=True)
rs['B9'] = '=COUNTIFS(Citas!$D$2:$D$1000,TODAY()+1,Citas!$J$2:$J$1000,"<>Cancelada")'

# ---------------------------------------------------------------- Instrucciones
ins = wb.create_sheet('Instrucciones', 0)
ins.column_dimensions['A'].width = 110
lines = [
    ('AGENDA DR. RIVAS · PLANTILLA', True),
    ('', False),
    ('1. Sube este archivo a Google Drive y ábrelo con Hojas de cálculo de Google (o impórtalo en una hoja nueva).', False),
    ('2. Completa la hoja "Configuración": correo de cada sede, horarios reales y, si quieres, un calendario por sede.', False),
    ('3. Sigue la guía LEEME.md (carpeta agenda-backend) para pegar el script y publicarlo. Al final te da una dirección web: se la pasas a quien programa el sitio.', False),
    ('4. A partir de ahí, cada cita que se agende en el sitio aparece como una fila nueva en la hoja "Citas".', False),
    ('', False),
    ('QUÉ HACE LA RECEPCIÓN', True),
    ('• Cambia la columna "Estado": Pendiente → Confirmada / Atendida / No asistió / Cancelada.', False),
    ('• Si pones "Cancelada", la hora vuelve a quedar libre en el sitio y el evento se borra del calendario.', False),
    ('• Puedes escribir en "Notas (uso interno)". No cambies las demás columnas.', False),
    ('• No borres filas de citas futuras: cámbiales el estado a "Cancelada".', False),
    ('', False),
    ('QUÉ HACE SOLO EL SISTEMA', True),
    ('• Guarda cada cita, evita horas repetidas y respeta el máximo de citas por día.', False),
    ('• Crea el evento en el Google Calendar de la sede e invita al paciente si dejó su correo (le queda en su calendario).', False),
    ('• Envía un aviso por correo a la sede con los datos de la cita.', False),
    ('• Un día antes, envía un recordatorio por correo al paciente.', False),
]
for i, (t, b) in enumerate(lines, start=1):
    c = ins.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if b:
        c.font = Font(bold=True, size=13 if i == 1 else 11)

wb.save(OUT)
print('ok', OUT)
